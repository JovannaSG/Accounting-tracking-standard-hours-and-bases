# core/audit.py
"""Глубокий аудит данных 1С:Фреш (Phase 2-3).

Checkpoint-состояние живёт в отдельной БД ``audit_history.db``
(``core.db.init_audit_db``), движок строго последователен по базе и
покрыт ``tests/test_audit.py``.

Запуск из CLI::

    python -m core.audit --base <имя> --start 2026-01-01 --end 2026-01-31 \\
        [--include-unmapped] [--pause 0.2] [--restart]

Отчёт пишется в ``exports/audit_*.xlsx`` (листы «Сводка», «Проблемы»,
«Gaps»).
"""
from __future__ import annotations

import argparse
import sys
import time
from datetime import datetime
from datetime import date
from typing import Any
from urllib.parse import quote

import openpyxl

from core import db
from core.api_client import OneCClient
from core.norms import load_doc_types
from core.reporters import BORDER, CENTER, HEADER_FILL, HEADER_FONT

_TIME = time  # патчится в тестах (паузы между запросами)

# ====================== ЧАНКИНГ ПО КАЛЕНДАРНЫМ МЕСЯЦАМ ======================


def _month_chunks(period_start: str, period_end: str) -> list[tuple[str, str]]:
    """Дробит период на календарные месяцы (включительно).

    ``("2026-01-15", "2026-03-10")`` ->
    ``[("2026-01-15","2026-01-31"), ("2026-02-01","2026-02-28"),
    ("2026-03-01","2026-03-10")]``. Точные границы reduce запросов
    1С:Фреш и дают естественные «фрагменты» для checkpoint'ов.
    """
    from calendar import monthrange

    s = date.fromisoformat(period_start)
    e = date.fromisoformat(period_end)
    if e < s:
        raise ValueError(
            f"Период некорректен: {period_start} позже {period_end}"
        )

    chunks: list[tuple[str, str]] = []
    cur = s
    while cur <= e:
        month_last = cur.replace(day=monthrange(cur.year, cur.month)[1])
        end = month_last if month_last < e else e
        chunks.append((cur.isoformat(), end.isoformat()))
        cur = date(cur.year + 1, 1, 1) if cur.month == 12 else date(
            cur.year, cur.month + 1, 1
        )
    return chunks


# =========================== УТИЛИТЫ ТРАНСПОРТА ==============================


def _amount(rec: dict) -> float:
    """Числовой итог документа; мусор/None → 0.0 (не роняем чанк)."""

    raw = rec.get("СуммаДокумента")
    try:
        return float(raw or 0.0)
    except (TypeError, ValueError):
        return 0.0


def _count_records(
    client: OneCClient,
    entity: str,
    chunk_start: str,
    chunk_end: str,
) -> tuple[int | None, bool]:
    """Лёгкий ``$top=1&$count=true`` с тем же фильтром по Date.

    Возвращает ``(total, available)``. Если 1С отвергла ``$count``
    (донор не публикует регистры/не знает оператора) — ``(None, False)``:
    аудит деградирует в export-only, без hard-fail.
    """
    start_safe = client._start_of_day(chunk_start)
    end_safe = client._end_of_day(chunk_end)
    ffilter = (
        f"Date ge datetime'{start_safe}' "
        f"and Date le datetime'{end_safe}'"
    )
    endpoint = (
        f"{client.base_url}/odata/standard.odata/{entity}"
        f"?{r'$format'}=json&{r'$filter'}={quote(ffilter, safe='')}"
        f"&{r'$top'}=1&{r'$count'}=true"
    )
    try:
        data = client._request_page(endpoint, {})
    except ValueError:
        return None, False
    count = data.get("@odata.count", data.get("odata.count"))
    if count is None:
        return None, False
    try:
        return int(count), True
    except (TypeError, ValueError):
        return None, False


# ============================= ДВИЖОК АУДИТА =================================


def _audit_entities(base: dict) -> list[str]:
    """Уникальные OData-сущности реестра в порядке doc_types.json.

    Учитывает ``bases.active_doc_types``: пустой список = все виды реестра
    (прецедент — core.fetch.fetch_documents).
    """
    registry = load_doc_types()
    keys = base.get("active_doc_types") or list(registry.keys())
    entities: list[str] = []
    for key in keys:
        spec = registry.get(key)
        if spec and spec.get("entity") and spec["entity"] not in entities:
            entities.append(spec["entity"])
    return entities


def run_audit(
    base: dict,
    period_start: str,
    period_end: str,
    pause: float = 0.2,
    include_unmapped: bool = False,
    client: OneCClient | None = None,
    restart: bool = False,
    entities: list[str] | None = None,
) -> dict:
    """Прогоняет аудит базы за период; резюмирует через checkpoint'ы.

    - (entity × месяц) как чанк: COMPLETED пропускаются, RUNNING/FAILED
      перезапускаются ровно с этой точки (без дублей сети);
    - каждый чанк: стрим ``fetch_documents_iter`` (память O(1)),
      агрегаты rows/сумма/страницы, затем ``$count``-gate;
    - деградация на минимальный ``$select`` (``client._degraded_entities``)
      фиксируется как data-poor;
    - ``include_unmapped`` — подсчёт нереестровых ``Document_*`` из
      ``$metadata`` (``$count``-only, с паузами).

    ``client``/``entities`` — тестовые интерфейсы: движок не строит сетевого
    клиента, если клиент уже передан; сущности — из реестра, а не сети.
    """
    base_url = str(base["url"]).rstrip("/")
    client = client or OneCClient(
        base_url, base.get("login") or "", base.get("password") or ""
    )

    init_audit_db = db.init_audit_db
    init_audit_db()
    conn = db._audit_conn()
    try:
        now = datetime.now().isoformat(timespec="seconds")
        conn.execute(
            "INSERT INTO audit_runs (base_url, period_start, period_end, "
            "include_unmapped, status, created_at) VALUES (?, ?, ?, ?, ?, ?)",
            (base_url, period_start, period_end, int(bool(include_unmapped)),
             db.AUDIT_RUN_PENDING, now),
        )
        run_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        conn.commit()
    finally:
        conn.close()

    summary: dict[str, Any] = {
        "run_id": run_id,
        "base_url": base_url,
        "base_name": base.get("name", ""),
        "period_start": period_start,
        "period_end": period_end,
        "chunks_total": 0,
        "chunks_completed": 0,
        "chunks_skipped": 0,
        "chunks_failed": 0,
        "count_mismatches": 0,
        "degraded_entities": [],
    }

    chunk_items = [
        (entity, chunk)
        for entity in (entities or _audit_entities(base))
        for chunk in _month_chunks(period_start, period_end)
    ]
    summary["chunks_total"] = len(chunk_items)

    try:
        for entity, (chunk_start, chunk_end) in chunk_items:
            status = _run_chunk(
                client, run_id, base_url, entity,
                chunk_start, chunk_end, pause, restart,
            )
            if status == "SKIPPED":
                summary["chunks_skipped"] += 1
            elif status == db.AUDIT_CP_COMPLETED:
                summary["chunks_completed"] += 1
            elif status == db.AUDIT_CP_FAILED:
                summary["chunks_failed"] += 1

        summary["degraded_entities"] = sorted(client._degraded_entities)

        if include_unmapped:
            gaps = _run_gap_counts(
                client, base_url, run_id,
                entities or _audit_entities(base), pause,
                period_start, period_end,
            )
            summary["gaps_count"] = len(gaps)

        summary["count_mismatches"] = _count_mismatches(run_id)

        run_status = (
            db.AUDIT_RUN_COMPLETED
            if summary["chunks_failed"] == 0
            else db.AUDIT_RUN_FAILED
        )
        _finalize_run(run_id, run_status)
    except Exception as e:
        _finalize_run(run_id, db.AUDIT_RUN_FAILED, error=str(e))
        raise

    summary["run_status"] = run_status
    return summary


def _run_chunk(
    client: OneCClient,
    run_id: int,
    base_url: str,
    entity: str,
    chunk_start: str,
    chunk_end: str,
    pause: float,
    restart: bool,
) -> str:
    """Аудит одного (entity × месяц); возвращает итоговый статус чанка."""

    conn = db._audit_conn()
    try:
        row = conn.execute(
            "SELECT " + ", ".join(db._AUDIT_CHECKPOINTS_COLUMNS)
            + " FROM audit_checkpoints WHERE base_url=? AND entity_name=?"
            + " AND chunk_start=? AND chunk_end=?",
            (base_url, entity, chunk_start, chunk_end),
        ).fetchone()
    finally:
        conn.close()

    if (row is not None and not restart
            and db._audit_row(row, db._AUDIT_CHECKPOINTS_COLUMNS)["status"]
            == db.AUDIT_CP_COMPLETED):
        return "SKIPPED"

    now = datetime.now().isoformat(timespec="seconds")
    conn = db._audit_conn()
    try:
        conn.execute(
            "INSERT INTO audit_checkpoints (run_id, base_url, entity_name, "
            "chunk_start, chunk_end, status, started_at, rows_processed) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, 0) "
            "ON CONFLICT(base_url, entity_name, chunk_start, chunk_end) "
            "DO UPDATE SET run_id=excluded.run_id, status='RUNNING', "
            "started_at=excluded.started_at, rows_processed=0, sum_amount=NULL, "
            "page_count=0, degraded=0, expected_count=NULL, "
            "count_available=0, count_mismatch=0, error=NULL, finished_at=NULL",
            (run_id, base_url, entity, chunk_start, chunk_end,
             db.AUDIT_CP_RUNNING, now),
        )
        conn.commit()
    finally:
        conn.close()

    rows = 0
    total = 0.0
    pages = 0

    def _on_page() -> None:
        nonlocal pages
        pages += 1

    try:
        for rec in client.fetch_documents_iter(
            entity, chunk_start, chunk_end, pause=pause, on_page=_on_page
        ):
            rows += 1
            total += _amount(rec)

        count, count_available = _count_records(
            client, entity, chunk_start, chunk_end
        )
        mismatch = (abs(count - rows) if count_available and count is not None
                    else 0)
        degraded = 1 if entity in client._degraded_entities else 0
        finished = datetime.now().isoformat(timespec="seconds")

        conn = db._audit_conn()
        try:
            conn.execute(
                "UPDATE audit_checkpoints SET status=?, rows_processed=?, "
                "sum_amount=?, page_count=?, degraded=?, expected_count=?, "
                "count_available=?, count_mismatch=?, finished_at=? "
                "WHERE base_url=? AND entity_name=? AND chunk_start=? "
                "AND chunk_end=?",
                (db.AUDIT_CP_COMPLETED, rows, round(total, 2), pages,
                 degraded, count, int(count_available), mismatch, finished,
                 base_url, entity, chunk_start, chunk_end),
            )
            conn.commit()
        finally:
            conn.close()
        return db.AUDIT_CP_COMPLETED
    except Exception as e:
        failed = datetime.now().isoformat(timespec="seconds")
        conn = db._audit_conn()
        try:
            conn.execute(
                "UPDATE audit_checkpoints SET status=?, error=?, finished_at=? "
                "WHERE base_url=? AND entity_name=? AND chunk_start=? "
                "AND chunk_end=?",
                (db.AUDIT_CP_FAILED, str(e)[:500], failed,
                 base_url, entity, chunk_start, chunk_end),
            )
            conn.commit()
        finally:
            conn.close()
        return db.AUDIT_CP_FAILED


def _run_gap_counts(
    client: OneCClient,
    base_url: str,
    run_id: int,
    registry_entities: list[str],
    pause: float,
    period_start: str,
    period_end: str,
) -> list[dict]:
    """Счётчики нереестровых ``Document_*`` (только --include-unmapped).

    Список сущностей берётся из ``$metadata``, нереестровые обходятся
    ``$count``-only с паузами; результат складывается в ``audit_gap_counts``
    и возвращается для отчёта.
    """
    registry_set = set(registry_entities)
    try:
        sets = client.fetch_metadata_entity_sets()
    except ValueError as e:
        raise ValueError(f"Не удалось получить $metadata для gaps: {e}") from e

    gaps: list[dict] = []
    now = datetime.now().isoformat(timespec="seconds")
    conn = db._audit_conn()
    try:
        for name in sorted(sets):
            if not name.startswith("Document_") or name in registry_set:
                continue
            if pause > 0:
                _TIME.sleep(pause)
            count, available = _count_records(
                client, name, period_start, period_end
            )
            conn.execute(
                "INSERT INTO audit_gap_counts (base_url, entity_name, run_id, "
                "doc_count, count_available, updated_at) "
                "VALUES (?, ?, ?, ?, ?, ?) "
                "ON CONFLICT(base_url, entity_name) DO UPDATE SET "
                "run_id=excluded.run_id, doc_count=excluded.doc_count, "
                "count_available=excluded.count_available, "
                "updated_at=excluded.updated_at",
                (base_url, name, run_id, count if count is not None else 0,
                 int(available), now),
            )
            gaps.append({
                "base_url": base_url,
                "entity_name": name,
                "run_id": run_id,
                "doc_count": count if count is not None else 0,
                "count_available": int(available),
                "updated_at": now,
            })
        conn.commit()
    finally:
        conn.close()
    return gaps


def _count_mismatches(run_id: int) -> int:
    conn = db._audit_conn()
    try:
        row = conn.execute(
            "SELECT COUNT(*) FROM audit_checkpoints WHERE run_id=? "
            "AND count_mismatch != 0",
            (run_id,),
        ).fetchone()
        return int(row[0])
    finally:
        conn.close()


def _finalize_run(run_id: int, status: str, error: str | None = None) -> None:
    now = datetime.now().isoformat(timespec="seconds")
    conn = db._audit_conn()
    try:
        if error is None:
            conn.execute(
                "UPDATE audit_runs SET status=?, finished_at=? WHERE id=?",
                (status, now, run_id),
            )
        else:
            conn.execute(
                "UPDATE audit_runs SET status=?, finished_at=? WHERE id=?",
                (db.AUDIT_RUN_FAILED, now, run_id),
            )
            conn.execute(
                "UPDATE audit_checkpoints SET status=?, error=? "
                "WHERE run_id=? AND status IN ('RUNNING','PENDING')",
                (db.AUDIT_CP_FAILED, error[:500], run_id),
            )
        conn.commit()
    finally:
        conn.close()


def _resolve_base(name_or_url: str) -> dict | None:
    """База из таблицы bases по имени (без учёта регистра), URL или id."""

    needle = str(name_or_url or "").strip()
    if not needle:
        return None
    for b in db.list_bases():
        if b["name"].strip().lower() == needle.lower():
            return b
    if needle.lower().startswith(("http://", "https://")):
        return db.get_base(needle)
    for b in db.list_bases():
        if str(b["id"]) == needle:
            return b
    return None


# ================================ ОТЧЁТ ======================================


def write_audit_report(
    summary: dict,
    out_dir: str = "exports",
) -> str:
    """Формирует ``exports/audit_*.xlsx``: «Сводка», «Проблемы», «Gaps».

    Каждая строка Сводки — checkpoint (entity × месяц). «Проблемы» — вход
    для ручной сверки: расхождения с ``$count`` (миграция данных), деградация
    (data-poor), ошибки выгрузки. Возвращает путь к файлу.
    """
    import os

    os.makedirs(out_dir, exist_ok=True)
    safe_name = "".join(
        c if (c.isalnum() or c in "-_.") else "_" for c in summary["base_name"]
    ).strip("_")
    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    path = os.path.join(
        out_dir,
        f"audit_{safe_name or 'base'}_{summary['period_start']}_"
        f"{summary['period_end']}_{ts}.xlsx",
    )
    checkpoints = db.list_audit_checkpoints(summary["run_id"])
    gaps = db.list_audit_gap_counts(summary["base_url"])

    workbook = openpyxl.Workbook()

    _write_summary_sheet(workbook, checkpoints, summary)
    _write_problems_sheet(workbook, checkpoints)
    _write_gaps_sheet(workbook, gaps)

    workbook.save(path)
    return path


def _style_headers(ws, headers: list[str]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"


def _cell(ws, row: int, col: int, value) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.border = BORDER
    return cell


def _write_summary_sheet(
    workbook: openpyxl.Workbook,
    checkpoints: list[dict],
    summary: dict,
) -> None:
    ws = workbook.active
    ws.title = "Сводка"
    headers = [
        "Сущность", "Чанк с", "Чанк по", "Строк", "Сумма документов",
        "Страниц", "Данных-бедно", "$count", "Ожидалось", "Расхождение",
        "Статус", "Ошибка",
    ]
    _style_headers(ws, headers)
    for row_i, cp in enumerate(checkpoints, start=2):
        values = [
            cp["entity_name"], cp["chunk_start"], cp["chunk_end"],
            cp["rows_processed"], cp["sum_amount"], cp["page_count"],
            "да" if cp["degraded"] else "",
            "да" if cp["count_available"] else "нет",
            cp["expected_count"], cp["count_mismatch"],
            cp["status"], cp["error"],
        ]
        for col_i, value in enumerate(values, start=1):
            _cell(ws, row_i, col_i, value)
    for col in (4, 5, 6, 9, 10):
        _fmt_column(ws, col, "0" if col in (4, 6, 9, 10) else "0.00")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["L"].width = 60


def _fmt_column(ws, col: int, fmt: str) -> None:
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=col).number_format = fmt


def _write_problems_sheet(
    workbook: openpyxl.Workbook,
    checkpoints: list[dict],
) -> None:
    ws = workbook.create_sheet("Проблемы")
    headers = ["Сущность", "Чанк с", "Чанк по", "Проблема", "Детали"]
    _style_headers(ws, headers)
    row_i = 2
    for cp in checkpoints:
        if cp["status"] == db.AUDIT_CP_FAILED:
            problem, detail = ("Ошибка выгрузки", cp["error"] or "")
        elif cp["count_mismatch"]:
            problem = "Расхождение с $count"
            detail = (
                f"ожидалось {cp['expected_count']}, загружено "
                f"{cp['rows_processed']}"
            )
        elif cp["degraded"]:
            problem = "Минимальный набор полей (data-poor)"
            detail = "400/500 на полном $select — деградация на 3 поля"
        else:
            continue
        for col_i, value in enumerate(
            [cp["entity_name"], cp["chunk_start"], cp["chunk_end"],
             problem, detail],
            start=1,
        ):
            _cell(ws, row_i, col_i, value)
        row_i += 1
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["E"].width = 70


def _write_gaps_sheet(
    workbook: openpyxl.Workbook,
    gaps: list[dict],
) -> None:
    ws = workbook.create_sheet("Gaps")
    headers = ["Сущность", "Количество документов", "$count доступен"]
    _style_headers(ws, headers)
    for row_i, gap in enumerate(gaps, start=2):
        for col_i, value in enumerate(
            [gap["entity_name"], gap["doc_count"],
             "да" if gap["count_available"] else "нет"],
            start=1,
        ):
            _cell(ws, row_i, col_i, value)
    _fmt_column(ws, 2, "0")
    ws.column_dimensions["A"].width = 46


# ================================== CLI ======================================


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m core.audit",
        description="Глубокий аудит данных баз 1С:Фреш (checkpoint-резюм).",
    )
    parser.add_argument("--base", required=True,
                        help="Имя (или URL/id) базы из таблицы bases")
    parser.add_argument("--start", required=True,
                        help="Начало периода, YYYY-MM-DD")
    parser.add_argument("--end", required=True,
                        help="Конец периода, YYYY-MM-DD")
    parser.add_argument("--include-unmapped", action="store_true",
                        help="Подсчитать нереестровые Document_* из $metadata")
    parser.add_argument("--pause", type=float, default=0.2,
                        help="Пауза между запросами, сек (по умолчанию 0.2)")
    parser.add_argument("--restart", action="store_true",
                        help="Перечитать COMPLETED-чанки заново")
    parser.add_argument("--out", default="exports",
                        help="Каталог отчёта (по умолчанию exports/)")
    args = parser.parse_args(argv)

    base = _resolve_base(args.base)
    if base is None:
        names = [b["name"] for b in db.list_bases()]
        print(
            f"База '{args.base}' не найдена. Доступные: {', '.join(names) or '-'}",
            file=sys.stderr,
        )
        return 2

    try:
        summary = run_audit(
            base,
            args.start,
            args.end,
            pause=args.pause,
            include_unmapped=args.include_unmapped,
            restart=args.restart,
        )
    except ValueError as e:
        print(f"Аудит не выполнен: {e}", file=sys.stderr)
        return 1

    path = write_audit_report(summary, out_dir=args.out)
    print(f"Отчёт: {path}")
    print(
        f"Чанков: {summary['chunks_total']} (готово "
        f"{summary['chunks_completed']}, ошибок {summary['chunks_failed']}), "
        f"расхождений с $count: {summary['count_mismatches']}"
    )
    if summary["degraded_entities"]:
        print("Data-poor (минимальный $select): "
              + ", ".join(summary["degraded_entities"]))
    if summary.get("gaps_count") is not None:
        print(f"Нереестровых Document_* подсчитано: {summary['gaps_count']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
