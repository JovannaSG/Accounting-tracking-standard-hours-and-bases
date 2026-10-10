# tests/test_audit.py
"""Phase 2 Deep Data Audit: движок run_audit (checkpoints, $count-gate, gaps).

Сеть не трогаем: подменяем session.get фейковым диспетчером, который отвечает
на страницы документов, "$count=true" и "$metadata" (см. доку, синхронно).
Проверяются: чанкинг, idempotent-резюм, расхождение с $count, деградация,
gaps, паузы, FAILED-обработка.
"""
import os
import types

import openpyxl
import pytest
import requests

from core import audit, db
from core.api_client import OneCClient

_BASE = {
    "name": "Тестовая база",
    "url": "https://audit.example/a/1",
    "login": "odata.user",
    "password": "secret",
}

_METADATA_XML = (
    '<?xml version="1.0" encoding="utf-8"?>'
    '<edmx:Edmx xmlns:edmx="http://docs.oasis-open.org/odata/ns/edmx" Version="1.0">'
    "<edmx:DataServices>"
    '<Schema xmlns="http://schemas.microsoft.com/ado/2008/09/edm" '
    'Namespace="StandardODATA">'
    '<EntityContainer Name="StandardOData">'
    '<EntitySet Name="Document_ПлатежноеПоручение" '
    'EntityType="StandardODATA.Document_ПлатежноеПоручение"/>'
    '<EntitySet Name="Document_НеРеестровое" '
    'EntityType="StandardODATA.Document_НеРеестровое"/>'
    "</EntityContainer></Schema></edmx:DataServices></edmx:Edmx>"
).encode("utf-8")


class _FakeResponse:
    """requests.Response-совместимость для session.get: json() и .content."""

    def __init__(self, payload=None, status=200, content=b""):
        self._payload = payload
        self._status = status
        self.content = content

    def raise_for_status(self):
        if self._status >= 400:
            response = requests.Response()
            response.status_code = self._status
            raise requests.exceptions.HTTPError(response=response)

    def json(self):
        return self._payload


def _row(sum_amount: float, ref: str = "r") -> dict:
    return {"Ref_Key": ref, "СуммаДокумента": sum_amount}


class _FakeData:
    """Сценарий выгрузки: сущность -> страницы/счётчик/ошибки.

    Параметры определяют поведение «1С»: какие сущности существуют,
    сколько строк в каждой, чему равен $count, какие сущности падают.
    """

    def __init__(self):
        self.requests_seen: list[str] = []
        self.entities: dict[str, dict] = {}  # entity -> {"rows": int,
        # "sum": float, "count": int|None, "fail": str|None, "degrade": bool}
        self.calls = 0

    def set_entity(self, name, *, rows=0, sum=0.0, count=None, fail=None,
                   degrade=False):
        self.entities[name] = {
            "rows": rows, "sum": sum, "count": count, "fail": fail,
            "degrade": degrade,
        }

    def fake_get(self, url, params=None, timeout=None):
        self.calls += 1
        self.requests_seen.append(url)
        if "/$metadata" in url:
            return _FakeResponse(content=_METADATA_XML)
        entity = "?" if "/odata/standard.odata/" not in url else url.split(
            "/odata/standard.odata/", 1)[1].split("?", 1)[0]
        spec = self.entities.get(entity)
        if spec is None:
            return _FakeResponse({"value": []}, 200)
        if "count=true" in url:
            if spec["count"] is None:
                return _FakeResponse({"value": []}, 400)
            return _FakeResponse(
                {"@odata.count": spec["count"], "value": [{"Ref_Key": "c"}]}
            )
        if spec["fail"]:
            return _FakeResponse(None, 400)
        if spec["degrade"]:
            from urllib.parse import parse_qs
            fields = parse_qs(url.split("?", 1)[1]).get("$select", [""])[0]
            if len(fields.split(",")) > 3:
                return _FakeResponse(None, 400)
        rows = [_row(spec["sum"] / spec["rows"] if spec["rows"] else 0.0,
                     f"{entity}#{i}") for i in range(spec["rows"])]
        return _FakeResponse({"value": rows})


class _FakeClient:
    """OneCClient с подменённым session.get (без старта реальной сети)."""

    def __init__(self, fake: _FakeData):
        self.client = OneCClient(_BASE["url"], _BASE["login"], _BASE["password"])
        self.client.session.get = fake.fake_get


def _checkpoint(base_url, entity, start, end) -> dict | None:
    conn = db._audit_conn()
    try:
        row = conn.execute(
            "SELECT " + ", ".join(db._AUDIT_CHECKPOINTS_COLUMNS)
            + " FROM audit_checkpoints WHERE base_url=? AND entity_name=?"
            + " AND chunk_start=? AND chunk_end=?",
            (base_url, entity, start, end),
        ).fetchone()
        return db._audit_row(row, db._AUDIT_CHECKPOINTS_COLUMNS) if row else None
    finally:
        conn.close()


# ============================= ЧАНКИНГ ======================================

def test_month_chunks_basic():
    assert audit._month_chunks("2026-01-15", "2026-03-10") == [
        ("2026-01-15", "2026-01-31"),
        ("2026-02-01", "2026-02-28"),
        ("2026-03-01", "2026-03-10"),
    ]


def test_month_chunks_single_month_and_day():
    assert audit._month_chunks("2026-01-01", "2026-01-31") == [
        ("2026-01-01", "2026-01-31")
    ]
    assert audit._month_chunks("2026-01-31", "2026-01-31") == [
        ("2026-01-31", "2026-01-31")
    ]


def test_month_chunks_across_year():
    assert audit._month_chunks("2025-12-15", "2026-01-10") == [
        ("2025-12-15", "2025-12-31"),
        ("2026-01-01", "2026-01-10"),
    ]


def test_month_chunks_rejects_inverted_period():
    with pytest.raises(ValueError):
        audit._month_chunks("2026-03-01", "2026-01-01")


# ============================= ДВИЖОК =======================================

def test_run_audit_completes_and_count_matches():
    fake = _FakeData()
    fake.set_entity("Document_A", rows=3, sum=300.0, count=3)
    fc = _FakeClient(fake)

    summary = audit.run_audit(
        _BASE, "2026-01-01", "2026-01-31",
        client=fc.client, pause=0.0,
        entities=["Document_A"],
    )

    assert summary["run_status"] == db.AUDIT_RUN_COMPLETED
    assert summary["chunks_total"] == 1
    assert summary["chunks_completed"] == 1
    assert summary["chunks_failed"] == 0
    assert summary["count_mismatches"] == 0

    cp = _checkpoint(_BASE["url"], "Document_A", "2026-01-01", "2026-01-31")
    assert cp["status"] == db.AUDIT_CP_COMPLETED
    assert cp["rows_processed"] == 3
    assert cp["sum_amount"] == 300.0
    assert cp["expected_count"] == 3
    assert cp["count_available"] == 1
    assert cp["count_mismatch"] == 0
    assert cp["degraded"] == 0
    assert cp["page_count"] == 1


def test_run_audit_multiple_entities_two_months():
    fake = _FakeData()
    fake.set_entity("Document_A", rows=2, sum=10.0, count=2)
    fake.set_entity("Document_B", rows=1, sum=5.0, count=1)
    fc = _FakeClient(fake)

    summary = audit.run_audit(
        _BASE, "2026-01-01", "2026-02-28",
        client=fc.client, pause=0.0,
        entities=["Document_A", "Document_B"],
    )

    assert summary["chunks_total"] == 4
    assert summary["chunks_completed"] == 4
    assert summary["run_status"] == db.AUDIT_RUN_COMPLETED
    # каждый чанк пишется своим checkpoint'ом
    cps = db.list_audit_checkpoints(summary["run_id"])
    assert len(cps) == 4
    assert {c["entity_name"] for c in cps} == {"Document_A", "Document_B"}


def test_count_mismatch_recorded():
    fake = _FakeData()
    fake.set_entity("Document_A", rows=1, sum=10.0, count=5)
    fc = _FakeClient(fake)

    summary = audit.run_audit(
        _BASE, "2026-01-01", "2026-01-31",
        client=fc.client, pause=0.0,
        entities=["Document_A"],
    )

    assert summary["count_mismatches"] == 1
    cp = _checkpoint(_BASE["url"], "Document_A", "2026-01-01", "2026-01-31")
    assert cp["count_mismatch"] == 4
    assert cp["expected_count"] == 5
    assert cp["rows_processed"] == 1


def test_count_unavailable_falls_back_to_export_only():
    fake = _FakeData()
    fake.set_entity("Document_A", rows=2, sum=20.0, count=None)  # $count 400
    fc = _FakeClient(fake)

    summary = audit.run_audit(
        _BASE, "2026-01-01", "2026-01-31",
        client=fc.client, pause=0.0,
        entities=["Document_A"],
    )

    cp = _checkpoint(_BASE["url"], "Document_A", "2026-01-01", "2026-01-31")
    assert cp["count_available"] == 0
    assert cp["count_mismatch"] == 0
    assert cp["rows_processed"] == 2
    # Запрос $count был сделан (не hard-fail)
    assert summary["run_status"] == db.AUDIT_RUN_COMPLETED


def test_resume_skips_completed_chunks():
    fake = _FakeData()
    fake.set_entity("Document_A", rows=3, sum=30.0, count=3)
    fc = _FakeClient(fake)

    first = audit.run_audit(
        _BASE, "2026-01-01", "2026-01-31",
        client=fc.client, pause=0.0,
        entities=["Document_A"],
    )
    n_before = fake.calls

    second = audit.run_audit(
        _BASE, "2026-01-01", "2026-01-31",
        client=fc.client, pause=0.0,
        entities=["Document_A"],
    )

    assert second["chunks_skipped"] == 1
    assert second["chunks_completed"] == 0
    # резюм не делает ни одного сетевого запроса
    assert fake.calls == n_before
    # Всего один checkpoint на этот (entity × chunk) навсегда
    assert len(db.list_audit_checkpoints(first["run_id"])) == 1


def test_restart_overrides_completed():
    fake = _FakeData()
    fake.set_entity("Document_A", rows=2, sum=20.0, count=2)
    fc = _FakeClient(fake)

    audit.run_audit(
        _BASE, "2026-01-01", "2026-01-31",
        client=fc.client, pause=0.0,
        entities=["Document_A"],
    )
    n_before = fake.calls
    summary = audit.run_audit(
        _BASE, "2026-01-01", "2026-01-31",
        client=fc.client, pause=0.0, restart=True,
        entities=["Document_A"],
    )

    assert summary["chunks_skipped"] == 0
    assert summary["chunks_completed"] == 1
    assert fake.calls > n_before  # перечитано заново


def test_degradation_marks_data_poor():
    fake = _FakeData()
    fake.set_entity("Document_C", rows=2, sum=10.0, count=2, degrade=True)
    fc = _FakeClient(fake)

    summary = audit.run_audit(
        _BASE, "2026-01-01", "2026-01-31",
        client=fc.client, pause=0.0,
        entities=["Document_C"],
    )

    assert "Document_C" in summary["degraded_entities"]
    cp = _checkpoint(_BASE["url"], "Document_C", "2026-01-01", "2026-01-31")
    assert cp["degraded"] == 1
    assert cp["count_mismatch"] == 0  # при деградации строки есть, gate честный


def test_entity_error_marks_checkpoint_and_run_failed():
    fake = _FakeData()
    fake.set_entity("Document_Err", rows=0, count=0, fail="always")
    fc = _FakeClient(fake)

    summary = audit.run_audit(
        _BASE, "2026-01-01", "2026-01-31",
        client=fc.client, pause=0.0,
        entities=["Document_Err"],
    )

    assert summary["run_status"] == db.AUDIT_RUN_FAILED
    assert summary["chunks_failed"] == 1
    cp = _checkpoint(_BASE["url"], "Document_Err", "2026-01-01", "2026-01-31")
    assert cp["status"] == db.AUDIT_CP_FAILED
    assert cp["error"]


# ============================= GAPS =========================================

def test_include_unmapped_counts_non_registry_documents():
    fake = _FakeData()
    fake.set_entity("Document_ПлатежноеПоручение", rows=1, sum=1.0, count=1)
    fake.set_entity("Document_НеРеестровое", rows=0, count=7)  # gap
    fc = _FakeClient(fake)

    summary = audit.run_audit(
        _BASE, "2026-01-01", "2026-01-31",
        client=fc.client, pause=0.0, include_unmapped=True,
        entities=["Document_ПлатежноеПоручение"],
    )

    assert summary["gaps_count"] == 1
    gaps = db.list_audit_gap_counts(_BASE["url"])
    assert [g["entity_name"] for g in gaps] == ["Document_НеРеестровое"]
    assert gaps[0]["doc_count"] == 7
    assert gaps[0]["count_available"] == 1

    # реестровая сущность в gaps не попадает
    assert all(g["entity_name"] != "Document_ПлатежноеПоручение" for g in gaps)


def test_gaps_without_flag_skips_metadata():
    fake = _FakeData()
    fake.set_entity("Document_ПлатежноеПоручение", rows=1, sum=1.0, count=1)
    fc = _FakeClient(fake)

    summary = audit.run_audit(
        _BASE, "2026-01-01", "2026-01-31",
        client=fc.client, pause=0.0, include_unmapped=False,
        entities=["Document_ПлатежноеПоручение"],
    )

    assert "gaps_count" not in summary
    assert db.list_audit_gap_counts(_BASE["url"]) == []
    assert not any("/$metadata" in u for u in fake.requests_seen)


# ============================= ПАУЗЫ ========================================

def test_pause_applied_between_count_and_gap_requests(monkeypatch):
    slept = []
    monkeypatch.setattr(
        audit, "_TIME", types.SimpleNamespace(sleep=lambda s: slept.append(s))
    )

    fake = _FakeData()
    fake.set_entity("Document_ПлатежноеПоручение", rows=0, count=1)
    fake.set_entity("Document_НеРеестровое", rows=0, count=3)
    fc = _FakeClient(fake)

    audit.run_audit(
        _BASE, "2026-01-01", "2026-01-31",
        client=fc.client, pause=0.5, include_unmapped=True,
        entities=["Document_ПлатежноеПоручение"],
    )

    assert slept, "с pause>0 ожидались паузы между запросами"
    assert all(round(s, 3) == 0.5 for s in slept)


# ============================= ОТЧЁТ ========================================

def test_write_audit_report_three_sheets(tmp_path):
    fake = _FakeData()
    fake.set_entity("Document_A", rows=1, sum=10.0, count=5)  # mismatch
    fake.set_entity("Document_B", rows=2, sum=20.0, count=2)  # ok
    fc = _FakeClient(fake)

    summary = audit.run_audit(
        _BASE, "2026-01-01", "2026-01-31",
        client=fc.client, pause=0.0,
        entities=["Document_A", "Document_B"],
    )
    path = audit.write_audit_report(summary, out_dir=str(tmp_path))

    assert os.path.exists(path)
    assert path.endswith(".xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Сводка", "Проблемы", "Gaps"]

    # Сводка: шапка + по строке на каждый чанк
    assert wb["Сводка"].max_row == len(
        db.list_audit_checkpoints(summary["run_id"])
    ) + 1

    # Проблемы: расхождение с $count по Document_A
    problems = wb["Проблемы"]
    assert problems.max_row >= 2
    problem_entities = [
        problems.cell(row=r, column=1).value
        for r in range(2, problems.max_row + 1)
    ]
    assert "Document_A" in problem_entities
    assert "Document_B" not in problem_entities


def test_main_unknown_base_returns_2(capsys):
    rc = audit.main([
        "--base", "ТакойБазыНет-12345",
        "--start", "2026-01-01", "--end", "2026-01-31",
    ])
    assert rc == 2
    err = capsys.readouterr().err
    assert "не найдена" in err


def test_resolve_base_by_name_and_url():
    base = db.insert_base(
        "Аудит Резолв", "https://resolve.example/a/1", "u", "p"
    )
    assert base is not None
    base_id = base["id"]
    by_name = audit._resolve_base("аудит резолв")
    by_url = audit._resolve_base("https://resolve.example/a/1")
    by_id = audit._resolve_base(str(base_id))
    assert by_name and by_name["id"] == base_id
    assert by_url and by_url["id"] == base_id
    assert by_id and by_id["id"] == base_id
    assert audit._resolve_base("") is None
